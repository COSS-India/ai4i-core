"""Compare the backfilled daily_usage rows (source = 'backfill') with Prometheus,
day by day, and write an Excel report (plus CSVs).

Read-only: it only runs SELECTs and Prometheus instant queries.

Only backfilled days are reported: the days in FROM..TO with backfill rows or
LLM usage in Prometheus. Days with neither are counted in the summary.

For every IST day in FROM..TO it asks Prometheus, at the day's end (next 00:00
IST), the same three questions the backfill asked: LLM chat requests, successful
LLM chat requests, and prompt / completion tokens, per institution, application
and service, with the dashboard's selectors (API-key traffic, real
institutions). Each (institution, application, service) group is rounded to a
whole number, as the backfill does, so the expected figures can be compared with
daily_usage exactly. Expected cost = rounded tokens x the service's current
price (mm_services, the PPU consumer's pricing rule).

Sheets:
  Summary            what was checked, how many days match
  Daily comparison   one row per backfilled day: Prometheus vs daily_usage for each figure
  Day x Institution  the same per institution
  Day x Service      tokens and requests per service
  Structural checks  days with usage have rows, no mixed sources, rows add up

Usage (from this folder):
  cp env.template .env          # then fill in the DB user and password
  python compare_backfill_vs_prometheus.py                    # full report
  python compare_backfill_vs_prometheus.py --prometheus-only  # no DB needed
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import httpx
from dotenv import dotenv_values

HERE = Path(__file__).resolve().parent
IST = timezone(timedelta(hours=5, minutes=30))
DAY = 86_400

# The backfill's selectors (services/pyjob/backfill/engine.py).
REQUESTS = "telemetry_obsv_requests_total"
TOKENS = "telemetry_obsv_llm_tokens_processed_sum"
CHAT_PATHS = r"/api/v1/chat(/completions)?"
BASE = 'tenant!="unknown",auth_type=~"api_key|"'
GROUP = "tenant_id, application_id, service_id"


# ── Config ──────────────────────────────────────────────────────────────────


def load_config() -> dict:
    cfg = {**dotenv_values(HERE / "env.template"), **dotenv_values(HERE / ".env")}
    return {k: (v or "").strip() for k, v in cfg.items()}


def days(first: date, last: date) -> list[date]:
    return [first + timedelta(n) for n in range((last - first).days + 1)]


def ist_midnight(d: date) -> datetime:
    return datetime.combine(d, time(), tzinfo=IST)


# ── Prometheus ──────────────────────────────────────────────────────────────


def day_sum(metric: str, group_by: str) -> str:
    """Counter growth over the 24 h before the evaluation time. Series that
    started in the window count their raw value, the rest their increase()."""
    return f"sum by({group_by}) (({metric} unless {metric} offset {DAY}s) or (increase({metric}[{DAY}s]) > 0))"


def queries(path_label: str) -> dict[str, str]:
    req = f'{REQUESTS}{{{BASE},{path_label}=~"{CHAT_PATHS}"}}'
    ok = f'{REQUESTS}{{{BASE},{path_label}=~"{CHAT_PATHS}",status_code=~"2.."}}'
    tok = f'{TOKENS}{{{BASE},token_type=~"prompt|completion"}}'
    return {
        "requests": day_sum(req, GROUP),
        "successful": day_sum(ok, GROUP),
        "tokens": day_sum(tok, f"{GROUP}, token_type"),
    }


def prom_query(client: httpx.Client, base: str, promql: str, at: datetime) -> list[dict]:
    r = client.get(f"{base}/api/v1/query", params={"query": promql, "time": at.timestamp()})
    r.raise_for_status()
    body = r.json()
    if body.get("status") != "success":
        raise RuntimeError(f"Prometheus query failed: {body}")
    return body["data"]["result"]


def label_int(value):
    return int(value) if value and value.isdigit() and int(value) > 0 else None


def whole(x: Decimal) -> Decimal:
    return x.quantize(Decimal(1), ROUND_HALF_UP)


@dataclass
class Group:
    """One (institution, application, service) group for a day."""
    raw: dict = field(default_factory=lambda: defaultdict(Decimal))   # unrounded Prometheus values

    def rounded(self) -> dict:
        req, ok = whole(self.raw["requests"]), whole(self.raw["successful"])
        return {"requests": req, "successful": req - max(Decimal(0), req - ok),
                "failed": max(Decimal(0), req - ok),
                "prompt": whole(self.raw["prompt"]), "completion": whole(self.raw["completion"])}


def prometheus_day(client, base, qs, d: date) -> dict[tuple, Group]:
    """{(tenant_id, application_id, service_id): Group} for day ``d``, keyed
    as the backfill keys its rows."""
    at = ist_midnight(d + timedelta(1))
    groups: dict[tuple, Group] = defaultdict(Group)
    for kind, promql in qs.items():
        for s in prom_query(client, base, promql, at):
            m = s["metric"]
            tenant = label_int(m.get("tenant_id"))
            if tenant is None:
                continue
            key = (tenant, label_int(m.get("application_id")), m.get("service_id") or None)
            name = m.get("token_type") if kind == "tokens" else kind
            groups[key].raw[name] += Decimal(str(s["value"][1]))
    # The backfill writes no row for a group with no requests and no tokens.
    return {k: g for k, g in groups.items()
            if not (whole(g.raw["requests"]) == 0 and whole(g.raw["prompt"]) + whole(g.raw["completion"]) == 0)}


# ── Database ────────────────────────────────────────────────────────────────


def connect(cfg):
    import psycopg2
    import psycopg2.extras
    return psycopg2.connect(
        host=cfg["DB_HOST"], port=int(cfg["DB_PORT"]), user=cfg["DB_USER"], password=cfg["DB_PASSWORD"],
        dbname=cfg["DB_NAME"], connect_timeout=15, cursor_factory=psycopg2.extras.RealDictCursor,
        options="-c default_transaction_read_only=on",
    )


def token_rates(unit_rate, cost_per_unit, unit_size, output_cost_per_unit):
    """The PPU consumer's LLM pricing (as the backfill): input = unit_rate, else
    cost_per_unit / unit_size; output = output_cost_per_unit / unit_size, else input."""
    inp = Decimal(0)
    if unit_rate:
        inp = Decimal(str(unit_rate))
    elif cost_per_unit and unit_size:
        inp = Decimal(str(cost_per_unit)) / int(unit_size)
    out = inp
    if output_cost_per_unit is not None and unit_size:
        out = Decimal(str(output_cost_per_unit)) / int(unit_size)
    return inp, out


def llm_prices(conn) -> dict[str, tuple[Decimal, Decimal]]:
    with conn.cursor() as cur:
        cur.execute("SELECT service_id, unit_rate, cost_per_unit, unit_size, output_cost_per_unit"
                    " FROM mm_services WHERE lower(task_type) = 'llm'")
        return {r["service_id"]: token_rates(r["unit_rate"], r["cost_per_unit"], r["unit_size"],
                                             r["output_cost_per_unit"]) for r in cur.fetchall()}


def db_rows(conn, first: date, last: date, source: str) -> list[dict]:
    with conn.cursor() as cur:
        cur.execute(
            """SELECT usage_date, tenant_id, max(tenant_name) AS tenant_name, service_id, count(*) AS rows,
                      sum(request_count) AS requests, sum(success_count) AS successful,
                      sum(failed_count) AS failed, sum(total_units - output_units - cached_input_units) AS prompt,
                      sum(cached_input_units) AS cached, sum(output_units) AS completion, sum(cost) AS cost
               FROM daily_usage
               WHERE source = %s AND usage_date BETWEEN %s AND %s AND (%s <> 'rollup' OR task_type = 'llm')
               GROUP BY usage_date, tenant_id, service_id""",
            (source, first, last, source),
        )
        return cur.fetchall()


def structural(conn, first: date, last: date, usage_days: set) -> list[dict]:
    checks = []
    with conn.cursor() as cur:
        def check(name, sql, args, ok, describe):
            cur.execute(sql, args)
            rows = cur.fetchall()
            checks.append({"Check": name, "Result": "PASS" if ok(rows) else "FAIL", "Details": describe(rows)})

        cur.execute("SELECT DISTINCT usage_date FROM daily_usage WHERE source = 'backfill'"
                    " AND usage_date BETWEEN %s AND %s", (first, last))
        have = {r["usage_date"] for r in cur.fetchall()}
        missing = sorted(d for d in usage_days if d not in have)
        extra = sorted(d for d in have if d not in usage_days)
        checks.append({"Check": "Every day with LLM usage in Prometheus has backfill rows (and no others)",
                       "Result": "PASS" if not missing and not extra else "FAIL",
                       "Details": f"{len(usage_days)} day(s) with usage, {len(have)} day(s) with rows"
                                  + (f"; missing: {', '.join(map(str, missing))}" if missing else "")
                                  + (f"; rows but no usage: {', '.join(map(str, extra))}" if extra else "")
                                  + ". Days with no usage get no rows, by design."})
        check("No day has both backfill and rollup rows",
              "SELECT usage_date FROM daily_usage WHERE usage_date BETWEEN %s AND %s"
              " GROUP BY usage_date HAVING count(DISTINCT source) > 1", (first, last), lambda r: not r,
              lambda r: "none" if not r else ", ".join(str(x["usage_date"]) for x in r))
        check("Backfill ends before usage_events begins",
              "SELECT (SELECT max(usage_date) FROM daily_usage WHERE source = 'backfill') AS last_backfill,"
              " (SELECT (min(occurred_at) AT TIME ZONE 'Asia/Kolkata')::date FROM usage_events) AS first_recorded",
              (), lambda r: r[0]["first_recorded"] is None or r[0]["last_backfill"] < r[0]["first_recorded"],
              lambda r: f"last backfilled day {r[0]['last_backfill']}, first recorded day {r[0]['first_recorded']}")
        check("success + failed = requests on every row",
              "SELECT count(*) AS n FROM daily_usage WHERE source = 'backfill'"
              " AND success_count + failed_count <> request_count", (), lambda r: r[0]["n"] == 0,
              lambda r: f"{r[0]['n']} bad row(s)")
        check("cost = input cost + cached cost + output cost on every row",
              "SELECT count(*) AS n FROM daily_usage WHERE source = 'backfill'"
              " AND cost <> input_units_cost + cached_input_units_cost + output_units_cost", (),
              lambda r: r[0]["n"] == 0, lambda r: f"{r[0]['n']} bad row(s)")
        check("Backfill rows outside FROM..TO",
              "SELECT count(*) AS n, min(usage_date) AS lo, max(usage_date) AS hi FROM daily_usage"
              " WHERE source = 'backfill' AND usage_date NOT BETWEEN %s AND %s", (first, last),
              lambda r: r[0]["n"] == 0,
              lambda r: "none" if not r[0]["n"] else f"{r[0]['n']} row(s), {r[0]['lo']}..{r[0]['hi']}")
        cur.execute(
            "SELECT count(*) AS rows, count(*) FILTER (WHERE tenant_name IS NULL) AS no_institution,"
            " count(*) FILTER (WHERE application_id IS NULL) AS no_application,"
            " count(*) FILTER (WHERE service_id IS NULL) AS unknown_service,"
            " count(*) FILTER (WHERE api_key_id IS NULL) AS no_single_key"
            " FROM daily_usage WHERE source = 'backfill' AND usage_date BETWEEN %s AND %s", (first, last))
        r = cur.fetchone()
        checks.append({"Check": "Rows the backfill couldn't fully resolve (info)", "Result": "INFO",
                       "Details": f"{r['rows']} rows: {r['no_institution']} without institution name, "
                                  f"{r['no_application']} without application, {r['unknown_service']} with a "
                                  f"service that isn't an LLM service in mm_services (cost 0), "
                                  f"{r['no_single_key']} without an API key (multi-key or no application)"})
    return checks




# ── Comparison ──────────────────────────────────────────────────────────────

FIGURES = ("requests", "successful", "failed", "prompt", "completion")


def prom_totals(groups: dict[tuple, Group], prices, by) -> dict:
    """Prometheus figures summed by ``by(key)``: raw, rounded per group (as the
    backfill), expected cost, and group count."""
    out = defaultdict(lambda: {"raw": defaultdict(Decimal), "rounded": defaultdict(Decimal),
                               "cost": Decimal(0), "groups": 0})
    for key, g in groups.items():
        t = out[by(key)]
        r = g.rounded()
        for f in FIGURES:
            t["rounded"][f] += r[f]
        for f in ("requests", "successful", "prompt", "completion"):
            t["raw"][f] += g.raw[f]
        t["raw"]["failed"] += max(Decimal(0), g.raw["requests"] - g.raw["successful"])
        rates = prices.get(key[2]) if prices is not None else None
        if rates:
            t["cost"] += r["prompt"] * rates[0] + r["completion"] * rates[1]
        t["groups"] += 1
    return out


ZERO = defaultdict(Decimal)   # daily_usage has no rows for the key: every figure is 0


def db_totals(rows, by) -> dict:
    out = defaultdict(lambda: defaultdict(Decimal))
    for r in rows:
        t = out[by(r)]
        for f in (*FIGURES, "cost", "rows", "cached"):
            t[f] += Decimal(str(r[f] or 0))
    return out


def num(x):
    if x is None:
        return None
    x = Decimal(x)
    return int(x) if x == x.to_integral_value() else float(round(x, 6))


def compare_row(label: dict, p, d, with_cost=True) -> dict:
    """One report row: per figure, Prometheus (raw), Prometheus (as backfilled),
    daily_usage, difference and status."""
    row = dict(label)
    statuses = []
    for f in FIGURES:
        exp = p["rounded"][f] if p else Decimal(0)
        got = d[f] if d is not None else None
        row[f"{f} | Prometheus raw"] = num(p["raw"][f]) if p else 0
        row[f"{f} | Prometheus (rounded per row)"] = num(exp)
        row[f"{f} | daily_usage"] = num(got)
        if got is None:
            row[f"{f} | diff"], row[f"{f} | status"] = None, "NO DB"
        else:
            diff = got - exp
            row[f"{f} | diff"] = num(diff)
            row[f"{f} | status"] = "MATCH" if diff == 0 else "MISMATCH"
            statuses.append(row[f"{f} | status"])
    if with_cost:
        exp = p["cost"] if p else Decimal(0)
        got = d["cost"] if d is not None else None
        row["cost | expected (tokens x current price)"] = num(exp)
        row["cost | daily_usage"] = num(got)
        if got is None:
            row["cost | diff"], row["cost | status"] = None, "NO DB"
        else:
            diff = got - exp
            row["cost | diff"] = num(diff)
            # daily_usage stores cost to 6 decimals per row.
            tol = Decimal("0.000001") * max(1, int(d["rows"]))
            row["cost | status"] = "MATCH" if abs(diff) <= tol else "MISMATCH"
            statuses.append(row["cost | status"])
    row["rows | Prometheus groups"] = p["groups"] if p else 0
    row["rows | daily_usage"] = num(d["rows"]) if d is not None else None
    row["Overall"] = ("NO DB" if d is None and statuses == [] else
                      "MATCH" if statuses and all(s == "MATCH" for s in statuses) else "MISMATCH")
    return row


def pct(a, b):
    a, b = Decimal(a or 0), Decimal(b or 0)
    if b == 0:
        return None if a == 0 else 100.0
    return float(round((a - b) / b * 100, 2))


# ── Output ──────────────────────────────────────────────────────────────────


def write_csv(path: Path, rows: list[dict]):
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    cols = list(rows[0].keys())
    for r in rows:
        cols += [k for k in r if k not in cols]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path: Path, sheets: dict[str, list[dict]]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fills = {"MATCH": "C6EFCE", "PASS": "C6EFCE", "MISMATCH": "FFC7CE", "FAIL": "FFC7CE",
             "NO DB": "FFEB9C", "INFO": "DDEBF7"}
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if not rows:
            ws.append(["(no rows)"])
            continue
        cols = list(rows[0].keys())
        for r in rows:
            cols += [k for k in r if k not in cols]
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9D9D9")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value in fills:
                    cell.fill = PatternFill("solid", fgColor=fills[cell.value])
        for i, c in enumerate(cols, 1):
            width = max(len(str(c)) * 0.55, *(len(str(r.get(c) or "")) for r in rows[:200]), 8)
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "B2" if name != "Summary" else None
    wb.save(path)


# ── Main ────────────────────────────────────────────────────────────────────


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--prometheus-only", action="store_true", help="skip the database (no daily_usage figures)")
    args = ap.parse_args()
    cfg = load_config()
    first, last = date.fromisoformat(cfg["BACKFILL_FROM"]), date.fromisoformat(cfg["BACKFILL_TO"])
    base = cfg["PROMETHEUS_URL"].rstrip("/")
    qs = queries(cfg["PROMETHEUS_API_PATH_LABEL"])
    out_dir = HERE / "report"
    out_dir.mkdir(exist_ok=True)

    conn = None if args.prometheus_only else connect(cfg)
    prices = llm_prices(conn) if conn else None

    with httpx.Client(timeout=float(cfg.get("PROMETHEUS_TIMEOUT_SECONDS") or 120)) as client:
        lowest = float(prom_query(client, base, "min(prometheus_tsdb_lowest_timestamp_seconds)",
                                  datetime.now(IST))[0]["value"][1])
        prom = {}
        for d in days(first, last):
            prom[d] = prometheus_day(client, base, qs, d)
            print(f"  Prometheus {d}: {len(prom[d])} groups", file=sys.stderr)

        backfill = db_rows(conn, first, last, "backfill") if conn else []
        by_day_db = db_totals(backfill, lambda r: r["usage_date"])
        by_inst_db = db_totals(backfill, lambda r: (r["usage_date"], r["tenant_id"]))
        names = {r["tenant_id"]: r["tenant_name"] for r in backfill if r["tenant_name"]}
        llm = set(prices) if prices is not None else None

        def svc(s):   # the backfill writes NULL for a service that isn't an LLM service in mm_services
            return s if llm is None or s in llm else None
        by_svc_db = db_totals(backfill, lambda r: (r["usage_date"], r["service_id"]))

        daily, inst, service = [], [], []
        backfilled = [d for d in days(first, last) if prom[d] or d in by_day_db]
        empty = len(days(first, last)) - len(backfilled)
        for d in backfilled:
            p_day = prom_totals(prom[d], prices, lambda k: d)
            daily.append(compare_row({"Date": d.isoformat()}, p_day.get(d), by_day_db.get(d, ZERO) if conn else None))
            p_inst = prom_totals(prom[d], prices, lambda k: k[0])
            for t in sorted(set(p_inst) | {k[1] for k in by_inst_db if k[0] == d}):
                inst.append(compare_row({"Date": d.isoformat(), "Institution id": t, "Institution": names.get(t)},
                                        p_inst.get(t), by_inst_db.get((d, t), ZERO) if conn else None))
            p_svc = prom_totals(prom[d], prices, lambda k: svc(k[2]))
            for s in sorted(set(p_svc) | {k[1] for k in by_svc_db if k[0] == d}, key=lambda x: x or ""):
                service.append(compare_row({"Date": d.isoformat(), "Service": s or "(not an LLM service / none)"},
                                           p_svc.get(s), by_svc_db.get((d, s), ZERO) if conn else None, with_cost=True))


    struct = structural(conn, first, last, {d for d, g in prom.items() if g}) if conn else []
    if conn:
        conn.close()

    # Summary
    def count(rows, status):
        return sum(1 for r in rows if r.get("Overall") == status)
    first_day_ok = lowest <= ist_midnight(first).timestamp()
    summary = [
        {"Item": "Range checked", "Value": f"{first} .. {last} (IST), {len(daily) + empty} days"},
        {"Item": "Backfilled days (reported below)", "Value": len(daily)},
        {"Item": "Days with no LLM usage (0 in Prometheus, no rows in daily_usage; not listed)", "Value": empty},
        {"Item": "Generated at", "Value": datetime.now(IST).strftime("%Y-%m-%d %H:%M IST")},
        {"Item": "Prometheus", "Value": f"{base} (path label: {cfg['PROMETHEUS_API_PATH_LABEL']})"},
        {"Item": "Prometheus oldest data", "Value": datetime.fromtimestamp(lowest, IST).strftime("%Y-%m-%d %H:%M IST")
         + (" - covers the first day from 00:00" if first_day_ok else " - AFTER the first day's 00:00: first day can't be checked")},
        {"Item": "Database", "Value": "not used (--prometheus-only)" if not conn and args.prometheus_only
         else f"{cfg['DB_HOST']}:{cfg['DB_PORT']}/{cfg['DB_NAME']} (read-only)"},
        {"Item": "Backfilled days: all figures match", "Value": count(daily, "MATCH")},
        {"Item": "Backfilled days: some figure mismatches", "Value": count(daily, "MISMATCH")},
        {"Item": "Day x Institution rows: match / mismatch", "Value": f"{count(inst, 'MATCH')} / {count(inst, 'MISMATCH')}"},
        {"Item": "Day x Service rows: match / mismatch", "Value": f"{count(service, 'MATCH')} / {count(service, 'MISMATCH')}"},
        {"Item": "Structural checks failed", "Value": sum(1 for c in struct if c["Result"] == "FAIL")},
        {"Item": "Totals | requests (Prometheus / daily_usage)",
         "Value": f"{sum(r['requests | Prometheus (rounded per row)'] for r in daily)} / "
                  f"{sum((r['requests | daily_usage'] or 0) for r in daily)}"},
        {"Item": "Totals | prompt tokens (Prometheus / daily_usage)",
         "Value": f"{sum(r['prompt | Prometheus (rounded per row)'] for r in daily)} / "
                  f"{sum((r['prompt | daily_usage'] or 0) for r in daily)}"},
        {"Item": "Totals | completion tokens (Prometheus / daily_usage)",
         "Value": f"{sum(r['completion | Prometheus (rounded per row)'] for r in daily)} / "
                  f"{sum((r['completion | daily_usage'] or 0) for r in daily)}"},
        {"Item": "Totals | cost (expected / daily_usage)",
         "Value": f"{round(sum(Decimal(str(r['cost | expected (tokens x current price)'] or 0)) for r in daily), 4)} / "
                  f"{round(sum(Decimal(str(r['cost | daily_usage'] or 0)) for r in daily), 4)}"},
        {"Item": "How to read",
         "Value": "'Prometheus (rounded per row)' is what the backfill should have written: each institution / "
                  "application / service group rounded to a whole number. It must equal daily_usage exactly "
                  "(MATCH). 'Prometheus raw' is the unrounded value; the small gap is rounding. Expected cost uses "
                  "today's mm_services prices, as the backfill did."},
    ]
    sheets = {"Summary": summary, "Daily comparison": daily, "Day x Institution": inst,
              "Day x Service": service, "Structural checks": struct}
    stamp = f"{first:%Y%m%d}_{last:%Y%m%d}"
    xlsx = out_dir / f"backfill_vs_prometheus_{stamp}.xlsx"
    write_xlsx(xlsx, sheets)
    for name, rows in sheets.items():
        write_csv(out_dir / f"{re.sub(r'[^a-z0-9]+', '_', name.lower()).strip('_')}.csv", rows)
    print(f"wrote {xlsx}")
    for r in summary[:-1]:
        print(f"  {r['Item']}: {r['Value']}")


if __name__ == "__main__":
    main()
